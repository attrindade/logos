"""Texto de anexos para triagem, LLM e sidecars (Etapa 4).

Toda função devolve texto vazio em falha, com uma linha de log: um anexo ruim nunca derruba a
nota. Dependências opcionais (pypdf, python-docx, trafilatura, ocrmypdf) são importadas dentro
das funções; Tesseract e ocrmypdf rodam como subprocesso com timeout.
"""
import base64
import logging
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import requests

from . import config

logger = logging.getLogger(__name__)

OCR_MIN_CHARS = 20  # abaixo disso a foto vai para descrição por modelo de visão
PDF_MIN_CHARS = 50  # abaixo disso o PDF é tratado como escaneado
OCR_TIMEOUT_S = 120
OCRMYPDF_TIMEOUT_S = 600
VISION_TIMEOUT_S = 600
LINK_TIMEOUT_S = 15
DEFUDDLE_TIMEOUT_S = 60
VISION_PROMPT = (
    "Descreva esta imagem em português do Brasil, em até 3 frases objetivas. "
    "Se houver texto legível, transcreva-o."
)


def _no_window() -> dict:
    return {"creationflags": subprocess.CREATE_NO_WINDOW} if sys.platform == "win32" else {}


def _tesseract_env() -> dict:
    """TESSDATA_PREFIX aponta para o tessdata com `por`; o PATH ganha a pasta do Tesseract, que o
    ocrmypdf procura pelo nome."""
    env = dict(os.environ)
    if config.TESSDATA_DIR and Path(config.TESSDATA_DIR).is_dir():
        env["TESSDATA_PREFIX"] = str(config.TESSDATA_DIR)
    if config.TESSERACT_CMD:
        env["PATH"] = str(Path(config.TESSERACT_CMD).parent) + os.pathsep + env.get("PATH", "")
    return env


def plain_text(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8-sig").strip()
    except (OSError, UnicodeDecodeError) as e:
        logger.warning(f"Falha ao ler {path.name}: {e}")
        return ""


# --- Imagens ---------------------------------------------------------------

def ocr_image(path: Path) -> str:
    if not config.TESSERACT_CMD:
        logger.info("Tesseract não encontrado: OCR pulado.")
        return ""
    try:
        result = subprocess.run(
            [config.TESSERACT_CMD, str(path), "stdout", "-l", config.OCR_LANGS],
            capture_output=True, timeout=OCR_TIMEOUT_S, env=_tesseract_env(), **_no_window(),
        )
    except (OSError, subprocess.TimeoutExpired) as e:
        logger.warning(f"OCR falhou em {path.name}: {e}")
        return ""
    if result.returncode != 0:
        logger.warning(f"OCR falhou em {path.name}: {result.stderr.decode(errors='replace').strip()[-300:]}")
        return ""
    return result.stdout.decode("utf-8", errors="replace").strip()


def describe_image(path: Path) -> str:
    """Descrição por modelo de visão do Ollama. Lento na CPU (medido: 78 s com carga do modelo)."""
    if not config.VISION_MODEL:
        return ""
    try:
        from .llm import _ollama_generate

        image = base64.b64encode(path.read_bytes()).decode("ascii")
        return _ollama_generate(VISION_PROMPT, model=config.VISION_MODEL, images=[image]).strip()
    except Exception as e:
        logger.warning(f"Descrição por visão falhou em {path.name}: {e}")
        return ""


def image_text(path: Path) -> str:
    """OCR primeiro; se vier (quase) vazio, descrição pelo modelo de visão."""
    text = ocr_image(path)
    if len(text) >= OCR_MIN_CHARS:
        return text
    description = describe_image(path)
    return f"(descrição automática) {description}" if description else text


# --- Documentos ------------------------------------------------------------

def _pypdf_text(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()


def _ocrmypdf_text(path: Path) -> str:
    if not config.TESSERACT_CMD:
        return ""
    with tempfile.TemporaryDirectory(prefix="logos-ocr-") as tmp:
        out = Path(tmp) / "ocr.pdf"
        cmd = [
            sys.executable, "-m", "ocrmypdf", "-l", config.OCR_LANGS, "--skip-text",
            "--optimize", "0", "--output-type", "pdf", str(path), str(out),
        ]
        try:
            result = subprocess.run(
                cmd, capture_output=True, timeout=OCRMYPDF_TIMEOUT_S, env=_tesseract_env(), **_no_window(),
            )
        except (OSError, subprocess.TimeoutExpired) as e:
            logger.warning(f"ocrmypdf falhou em {path.name}: {e}")
            return ""
        if result.returncode != 0 or not out.exists():
            logger.warning(f"ocrmypdf falhou em {path.name}: {result.stderr.decode(errors='replace').strip()[-300:]}")
            return ""
        try:
            return _pypdf_text(out)
        except Exception as e:
            logger.warning(f"Falha ao ler o PDF do ocrmypdf de {path.name}: {e}")
            return ""


def pdf_text(path: Path) -> str:
    """Camada de texto com pypdf; sem ela (PDF escaneado), ocrmypdf numa cópia temporária."""
    try:
        text = _pypdf_text(path)
    except Exception as e:
        logger.warning(f"Falha ao extrair texto do PDF {path.name}: {e}")
        text = ""
    if len(text) >= PDF_MIN_CHARS:
        return text
    ocr = _ocrmypdf_text(path)
    return ocr if len(ocr) > len(text) else text


def docx_text(path: Path) -> str:
    try:
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    except Exception as e:
        logger.warning(f"Falha ao extrair texto do .docx {path.name}: {e}")
        return ""


def document_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return pdf_text(path)
    if ext == ".docx":
        return docx_text(path)
    if ext in (".md", ".txt"):
        return plain_text(path)
    return ""


# --- Links -----------------------------------------------------------------

def _defuddle(url: str) -> str:
    exe = shutil.which("defuddle")
    if not exe:
        return ""
    try:
        result = subprocess.run(
            [exe, "parse", url, "--md"], capture_output=True, timeout=DEFUDDLE_TIMEOUT_S, **_no_window(),
        )
    except (OSError, subprocess.TimeoutExpired) as e:
        logger.warning(f"defuddle falhou em {url}: {e}")
        return ""
    return result.stdout.decode("utf-8", errors="replace").strip() if result.returncode == 0 else ""


def link_text(url: str) -> tuple[str, str]:
    """(título, artigo em Markdown). `defuddle` se estiver no PATH; senão `trafilatura`."""
    text = _defuddle(url)
    if text:
        first = text.splitlines()[0].lstrip("# ").strip() if text.startswith("#") else ""
        return first, text
    try:
        import trafilatura

        resp = requests.get(url, timeout=LINK_TIMEOUT_S, headers={"User-Agent": "Mozilla/5.0 (Logos)"})
        resp.raise_for_status()
        article = trafilatura.extract(resp.text, url=url, output_format="markdown") or ""
        metadata = trafilatura.extract_metadata(resp.text)
        title = metadata.title if metadata is not None and metadata.title else ""
        return title.strip(), article.strip()
    except Exception as e:
        logger.warning(f"Falha ao extrair o link {url}: {e}")
        return "", ""
